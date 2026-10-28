# -*- coding: utf-8 -*-
"""AI 历史人物模拟器 —— 研究报告型 PPT（约 20 页）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
import os

OUT = "/Users/rogan09/Documents/Codex/2026-10-09/new-chat/outputs/History_AI_Research_Report.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x3A, 0x6E, 0xA5)
TEAL = RGBColor(0x2F, 0x6F, 0x5E)
AMBER = RGBColor(0xC8, 0x62, 0x2A)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x6B, 0x76, 0x86)
LINE = RGBColor(0xD8, 0xDE, 0xE7)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CJK = "PingFang SC"
LATIN = "Arial"

SW, SH = 13.333, 7.5
M = 0.85

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
_blank = prs.slide_layouts[6]


def _apply_font(run, size, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = LATIN
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rpr.append(ea)
    ea.set("typeface", CJK)


def box(slide, x, y, w, h, text, size=16, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        _apply_font(r, size, bold, color, italic)
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.1)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def new_slide():
    return prs.slides.add_slide(_blank)


def page_number(slide, n):
    box(slide, SW - M - 0.9, SH - 0.55, 0.9, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def header(slide, title, kicker=None, n=None):
    rect(slide, 0, 0, 0.14, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
    y = 0.55
    if kicker:
        box(slide, M, y, 11, 0.28, kicker.upper(), size=10.5, bold=True, color=BLUE)
        y += 0.32
    box(slide, M, y, 11.6, 0.85, title, size=30, bold=True, color=NAVY, spacing=0.95)
    rect(slide, M, y + 0.9, 0.95, 0.04, BLUE, shape=MSO_SHAPE.RECTANGLE)
    if n:
        page_number(slide, n)
    return y + 1.25


def bullets(slide, x, y, w, items, size=15.5, gap=0.16, color=INK):
    cy = y
    for it in items:
        head, sub = (it if isinstance(it, tuple) else (it, None))
        rect(slide, x, cy + 0.08, 0.085, 0.085, BLUE, shape=MSO_SHAPE.OVAL)
        box(slide, x + 0.26, cy, w - 0.26, 0.4, head, size=size, color=color, spacing=1.1)
        cy += 0.34
        if sub:
            box(slide, x + 0.26, cy + 0.02, w - 0.26, 0.55, sub, size=size - 3,
                color=MUTED, spacing=1.2)
            cy += 0.42
        cy += gap + 0.16
    return cy


def source(slide, text):
    box(slide, M, SH - 0.52, 11.2, 0.3, text, size=9, color=MUTED)


# ---------------------------------------------------------------- 1 封面
s = new_slide()
rect(s, 0, 0, SW, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.2, SH, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 1.75, 11, 0.35, "AI IN EDUCATION  ·  RESEARCH REPORT", size=12, bold=True,
    color=RGBColor(0x8F, 0xB6, 0xE0))
box(s, M, 2.25, 11.6, 1.3, "当 AI 比学生更会「给答案」", size=46, bold=True, color=WHITE, spacing=0.95)
box(s, M, 3.55, 10.8, 1.0, "用「不可靠的历史叙述者」训练历史思维 —— 一项设计与评估研究",
    size=19, color=RGBColor(0xC7, 0xD3, 0xE4), spacing=1.2)
rect(s, M, 4.75, 1.2, 0.04, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 4.98, 11, 0.35, "[小组成员]   ·   AI in Education   ·   2026", size=12.5,
    color=RGBColor(0x9F, 0xAE, 0xC4))
page_number(s, 1)

# ---------------------------------------------------------------- 2 研究问题
s = new_slide(); y = header(s, "我们想回答三个问题", "RESEARCH QUESTIONS", 2)
qs = [
    ("Q1", "AI 进入学习之后，学生获得知识的路径发生了什么变化？", "从「找答案」到「交给 AI 给答案」。"),
    ("Q2", "历史思维能否被 AI 训练，而不是被 AI 替代？", "溯源、语境化、交叉验证——恰是最容易被代劳的部分。"),
    ("Q3", "这种能力如何评估？评分可信吗？", "需要可测量的指标 + 可验证的评分信度。"),
]
cy = y + 0.1
for q, head, sub in qs:
    rect(s, M, cy, 11.6, 1.15, WHITE, line=LINE)
    rect(s, M, cy, 0.07, 1.15, BLUE, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.35, cy + 0.22, 0.7, 0.5, q, size=18, bold=True, color=BLUE)
    box(s, M + 1.15, cy + 0.2, 9.9, 0.4, head, size=15.5, bold=True, color=NAVY)
    box(s, M + 1.15, cy + 0.66, 9.9, 0.35, sub, size=12.5, color=MUTED)
    cy += 1.35
box(s, M, cy + 0.1, 11.6, 0.4, "我们不问「AI 能不能教」，而问「AI 会不会让学生不再需要思考」。",
    size=14, italic=True, color=BLUE)

# ---------------------------------------------------------------- 3 背景①
s = new_slide(); y = header(s, "用 AI 做作业，正在从「少数人」变成「多数人」", "背景 ①", 3)
cd = CategoryChartData()
cd.categories = ["2025 年 5 月", "2025 年 12 月"]
cd.add_series("使用 AI 做作业的比例", (48, 62))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(y + 0.15),
                        Inches(6.6), Inches(3.9), cd)
ch = gf.chart
ch.has_legend = False
ch.font.size = Pt(12)
ch.font.color.rgb = INK
plot = ch.plots[0]
plot.gap_width = 90
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '0"%"'; dl.number_format_is_linked = False
dl.font.size = Pt(16); dl.font.bold = True; dl.position = XL_LABEL_POSITION.OUTSIDE_END
ser = plot.series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = BLUE
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 80
ch.value_axis.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.size = Pt(12)

rect(s, 7.85, y + 0.25, 4.6, 1.35, WHITE, line=LINE)
box(s, 8.1, y + 0.42, 4.1, 0.55, "+14 个百分点", size=30, bold=True, color=AMBER)
box(s, 8.1, y + 1.02, 4.1, 0.4, "7 个月内（n = 1,214 名青少年）", size=12, color=MUTED)
bullets(s, 7.9, y + 1.95, 4.55, [
    "45% 的学生在学校用过 AI（HEPI, 2025）",
    "85% 使用 AI 的青少年把它用于学业（Common Sense Media）",
], size=12.5, gap=0.24)
source(s, "Sources: RAND (2026-03, n=1,214), HEPI Student Generative AI Survey (2025-02), Common Sense Media (2025)")

# ---------------------------------------------------------------- 4 背景②
s = new_slide(); y = header(s, "但用途高度集中在「把问题交给 AI」", "背景 ②", 4)
cd = CategoryChartData()
cd.categories = ["获得更好的解释", "发散想法", "查事实"]
cd.add_series("使用比例", (38, 35, 33))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(M), Inches(y + 0.2),
                        Inches(7.6), Inches(3.5), cd)
ch = gf.chart
ch.has_legend = False
ch.font.size = Pt(12); ch.font.color.rgb = INK
plot = ch.plots[0]
plot.gap_width = 70
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '0"%"'; dl.number_format_is_linked = False
dl.font.size = Pt(15); dl.font.bold = True
ser = plot.series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = TEAL
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 50
ch.category_axis.tick_labels.font.size = Pt(13)
ch.category_axis.reverse_order = True
rect(s, 8.75, y + 0.35, 3.7, 3.1, LIGHT)
box(s, 9.05, y + 0.65, 3.1, 0.4, "读数", size=12, bold=True, color=BLUE)
box(s, 9.05, y + 1.05, 3.1, 2.2,
    "三种高频用途，本质都是\n「让 AI 替我想」。\n\n几乎没有一项是\n「用 AI 检验我的判断」。",
    size=13.5, color=INK, spacing=1.45)
source(s, "Source: CSBA Blog 引用的学生调查（2026-04）")

# ---------------------------------------------------------------- 5 背景③
s = new_slide(); y = header(s, "教师的担忧：不是「用不用」，而是「还思不思考」", "背景 ③", 5)
cards = [
    ("54%", "教师认为 AI 让学生更难学会批判性思考", "美国教师民调（2026-06）"),
    ("43.2%", "教师担忧学生使用生成式 AI", "同比上升 7.5 个百分点（National Literacy Trust, 2025-08）"),
    ("66%", "英格兰中学教师表达担忧", "针对数千名教师的调查（2026）"),
]
for i, (num, desc, src) in enumerate(cards):
    x = M + i * 3.95
    rect(s, x, y + 0.1, 3.6, 2.75, WHITE, line=LINE)
    rect(s, x, y + 0.1, 3.6, 0.07, [BLUE, AMBER, TEAL][i], shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.3, y + 0.45, 3.0, 0.7, num, size=34, bold=True, color=NAVY)
    box(s, x + 0.3, y + 1.25, 3.0, 0.9, desc, size=12.5, color=INK, spacing=1.3)
    box(s, x + 0.3, y + 2.2, 3.0, 0.45, src, size=9.5, color=MUTED, spacing=1.2)
box(s, M, y + 3.05, 11.6, 0.5, "担忧的核心不是工具本身，而是「思考过程被跳过」。", size=17, bold=True, color=NAVY)
source(s, "Sources: Central Florida Public Media (2026-06), National Literacy Trust (2025-08), England teacher survey (2026)")

# ---------------------------------------------------------------- 6 痛点
s = new_slide(); y = header(s, "历史学科最怕的，正是最容易被代劳的那部分", "痛点", 6)
chain = ["AI 直接给出结论", "学生跳过求证过程", "溯源 / 语境化 / 交叉验证 被架空", "只剩「知道结论」"]
cw = 2.55
for i, t in enumerate(chain):
    x = M + i * (cw + 0.3)
    rect(s, x, y + 0.35, cw, 1.15, WHITE, line=LINE)
    box(s, x + 0.15, y + 0.45, cw - 0.3, 0.95, t, size=12.5, color=INK,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.25)
    if i < len(chain) - 1:
        box(s, x + cw, y + 0.75, 0.3, 0.3, "→", size=16, color=BLUE, align=PP_ALIGN.CENTER)
bullets(s, M, y + 1.85, 11.4, [
    ("历史思维不是「知道发生了什么」，而是「凭什么相信它发生过」。",
     "史料溯源、语境化、交叉验证——这三件事无法靠听讲获得，只能练。"),
    ("AI 越流畅，学生越没有理由去查证。",
     "这正是我们要设计的「逆操作」：不是让 AI 更好用，而是让它必须被怀疑。"),
], size=15, gap=0.2)
box(s, M, y + 3.55, 11.6, 0.5, "结论：历史学科需要的不是「更聪明的 AI」，而是「值得被质疑的 AI」。",
    size=16, bold=True, color=BLUE)

# ---------------------------------------------------------------- 7 现有做法
s = new_slide(); y = header(s, "现有 AI + 教育的五种主流做法", "LITERATURE", 7)
rows = [
    ("AI 答疑 / 智能导师", "对话式讲解、逐步提示", "降低理解门槛", "缩短学生与答案的距离"),
    ("作文自动批改", "评分 + 修改建议", "减轻教师负担", "同上"),
    ("题库 / 练习生成", "按知识点生成题目", "规模化练习", "同上"),
    ("语言翻译与陪练", "即时翻译、口语对话", "跨越语言障碍", "同上"),
    ("角色扮演", "与历史人物 / 情境对话", "提升沉浸感", "多为「聊天」，少为「求证」"),
]
rect(s, M, y, 11.6, 0.4, NAVY, shape=MSO_SHAPE.RECTANGLE)
for i, h in enumerate(["做法", "典型形态", "解决什么", "共同点"]):
    box(s, M + 0.18 + [0, 2.9, 6.0, 9.0][i], y + 0.1, [2.7, 3.0, 2.9, 2.5][i], 0.3,
        h, size=11.5, bold=True, color=WHITE)
cy = y + 0.55
for k, (a, b, c, d) in enumerate(rows):
    if k % 2 == 0:
        rect(s, M, cy - 0.06, 11.6, 0.72, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.18, cy, 2.7, 0.6, a, size=12.5, bold=True, color=NAVY)
    box(s, M + 3.08, cy, 2.9, 0.6, b, size=11.5, color=MUTED)
    box(s, M + 6.18, cy, 2.8, 0.6, c, size=11.5, color=INK)
    box(s, M + 9.18, cy, 2.4, 0.6, d, size=11.5, color=BLUE if k == 4 else MUTED)
    cy += 0.78
box(s, M, cy + 0.12, 11.6, 0.45, "五种做法都在做同一件事：缩短学生与答案之间的距离。",
    size=15, italic=True, color=NAVY)

# ---------------------------------------------------------------- 8 研究缺口
s = new_slide(); y = header(s, "研究缺口：把「AI 会错」当成教学材料", "GAP", 8)
gx, gy, gw, gh = M + 0.9, y + 0.15, 6.4, 3.6
rect(s, gx, gy, gw, gh, WHITE, line=LINE, shape=MSO_SHAPE.RECTANGLE)
rect(s, gx + gw / 2, gy, 0.02, gh, LINE, shape=MSO_SHAPE.RECTANGLE)
rect(s, gx, gy + gh / 2, gw, 0.02, LINE, shape=MSO_SHAPE.RECTANGLE)
box(s, gx, gy + gh + 0.12, gw, 0.3, "AI 的角色：给答案  ←→  被审视", size=11, color=MUTED, align=PP_ALIGN.CENTER)
box(s, gx - 1.35, gy + gh / 2 - 0.2, 1.25, 0.4, "通用", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
box(s, gx - 1.35, gy + gh - 0.2, 1.25, 0.4, "历史 / 史料", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
rect(s, gx, gy, gw / 2, gh / 2, LIGHT, shape=MSO_SHAPE.RECTANGLE)
box(s, gx + 0.22, gy + 0.3, 2.8, 0.8, "AI 答疑 / 批改\n（现有工作主流）", size=11.5, color=MUTED, spacing=1.35)
box(s, gx + gw / 2 + 0.22, gy + 0.3, 2.8, 0.6, "较少见", size=11.5, color=MUTED)
rect(s, gx, gy + gh / 2, gw / 2, gh / 2, LIGHT, shape=MSO_SHAPE.RECTANGLE)
box(s, gx + 0.22, gy + gh / 2 + 0.3, 2.8, 0.8, "历史人物聊天\n（现有工作主流）", size=11.5, color=MUTED, spacing=1.35)
rect(s, gx + gw / 2 + 0.22, gy + gh / 2 + 0.22, gw / 2 - 0.44, gh / 2 - 0.44, NAVY)
box(s, gx + gw / 2 + 0.42, gy + gh / 2 + 0.42, 2.6, 1.0, "本项目\nAI 被审视 × 史料", size=12.5, bold=True, color=WHITE, spacing=1.35)
bullets(s, 8.5, y + 0.3, 4.0, [
    "角色扮演多用于语言学习",
    "历史思维评估少见于交互系统",
    "把「模型不可靠」当教材的几乎没有",
], size=12, gap=0.1)
source(s, "定位基于对现有 AI+教育 工作的归纳，非穷尽性综述。")

# ---------------------------------------------------------------- 9 方案
s = new_slide(); y = header(s, "我们的方案：让 AI 成为一个「值得被质疑」的当事人", "OUR APPROACH", 9)
box(s, M, y, 11.6, 0.75,
    "AI 扮演有立场、可能出错的历史人物；学生必须用史料交叉验证，再为自己的决策辩护。",
    size=17, bold=True, color=NAVY, spacing=1.25)
steps = [("追问", "向人物提问、质疑"), ("核对史料", "判断说法与史料是否相符"),
         ("做抉择", "在历史节点上选择"), ("提交论证", "写出选择与依据"), ("反馈与画像", "五维评分 + 自评对比")]
bw, gap = 2.14, 0.24
for i, (t, d) in enumerate(steps):
    x = M + i * (bw + gap)
    rect(s, x, y + 1.15, bw, 1.5, WHITE, line=LINE)
    rect(s, x, y + 1.15, bw, 0.07, TEAL, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.12, y + 1.35, bw - 0.24, 0.35, t, size=13.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    box(s, x + 0.12, y + 1.78, bw - 0.24, 0.7, d, size=10.5, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.15)
    if i < len(steps) - 1:
        box(s, x + bw, y + 1.72, gap, 0.3, "→", size=15, color=BLUE, align=PP_ALIGN.CENTER)
box(s, M, y + 2.95, 11.6, 0.5, "关键区别：AI 不是「答案来源」，而是「需要被检验的信源」。",
    size=15, italic=True, color=BLUE)

# ---------------------------------------------------------------- 10 架构
s = new_slide(); y = header(s, "系统由四个模块组成", "SYSTEM", 10)
mods = [("场景引擎", "设定历史时刻与决策点"), ("史料库", "一手/二手史料，带出处"),
        ("角色智能体", "扮演人物，受史料约束"), ("诊断与教师端", "评分、画像、学情分析")]
for i, (t, d) in enumerate(mods):
    x = M + i * 2.98
    rect(s, x, y + 0.1, 2.72, 1.45, WHITE, line=LINE)
    rect(s, x, y + 0.1, 2.72, 0.07, BLUE, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.22, y + 0.33, 2.3, 0.35, t, size=14, bold=True, color=NAVY)
    box(s, x + 0.22, y + 0.8, 2.3, 0.6, d, size=11, color=MUTED, spacing=1.15)
box(s, M, y + 1.85, 11.6, 1.2,
    "实现上只用「预训练模型 API + 提示工程 + 小型史料库」，不训练、不微调任何模型。\n"
    "这是一个自觉的范围决策：把有限时间放在教学设计与评估上。",
    size=14, color=INK, spacing=1.4)
box(s, M, y + 3.15, 11.6, 0.4, "[ 此处放系统截图 ]", size=12, color=MUTED, italic=True)

# ---------------------------------------------------------------- 11 学生使用
s = new_slide(); y = header(s, "学生怎么用：四步走完一个学习闭环", "USAGE · 学生", 11)
student_steps = [("1", "选人物", "从人物卡片进入对话"), ("2", "追问", "质疑它的说法与依据"),
                 ("3", "核对史料", "判断说法与史料是否一致"), ("4", "抉择 + 论证", "选择并写出理由，再做自评")]
for i, (num, t, d) in enumerate(student_steps):
    x = M + i * 2.98
    rect(s, x, y + 0.25, 2.72, 2.4, WHITE, line=LINE)
    box(s, x + 0.25, y + 0.45, 0.6, 0.6, num, size=22, bold=True, color=BLUE)
    box(s, x + 0.25, y + 1.15, 2.2, 0.35, t, size=14, bold=True, color=NAVY)
    box(s, x + 0.25, y + 1.6, 2.2, 0.8, d, size=11, color=MUTED, spacing=1.2)
box(s, M, y + 2.9, 11.6, 0.4, "[ 此处放学生端截图：对话 / 核对史料 / 抉择评分 ]", size=12, color=MUTED, italic=True)

# ---------------------------------------------------------------- 12 教师使用
s = new_slide(); y = header(s, "教师怎么用：三步把人物「开箱即用」地建起来", "USAGE · 教师", 12)
teacher_steps = [("1", "生成限制词", "填几句话即可生成；也可拖入现成文档"),
                 ("2", "编辑史料与抉择", "逐条增删改，支持拖拽排序"),
                 ("3", "发布", "学生端立即生效")]
for i, (num, t, d) in enumerate(teacher_steps):
    x = M + i * 3.95
    rect(s, x, y + 0.3, 3.6, 2.5, WHITE, line=LINE)
    box(s, x + 0.3, y + 0.5, 0.6, 0.6, num, size=22, bold=True, color=TEAL)
    box(s, x + 0.3, y + 1.2, 3.0, 0.35, t, size=14, bold=True, color=NAVY)
    box(s, x + 0.3, y + 1.68, 3.0, 0.8, d, size=11.5, color=MUTED, spacing=1.2)
box(s, M, y + 3.0, 11.6, 0.4, "[ 此处放教师端截图：向导 / 拖拽上传 / 抉择编辑 ]", size=12, color=MUTED, italic=True)

# ---------------------------------------------------------------- 13 评估设计
s = new_slide(); y = header(s, "评估：把「历史思维」拆成五个可观察的维度", "ASSESSMENT", 13)
dims = [("史料溯源", "是否追问或标注信息的来源与立场"),
        ("语境化", "是否把决策放回当时的条件中"),
        ("交叉验证", "是否对照多个来源而非只听一家"),
        ("论证", "是否有「论点—证据—推理」的链条"),
        ("历史共情", "是否理解当时人的处境与限制")]
for i, (t, d) in enumerate(dims):
    row = i // 2
    col = i % 2
    x = M + col * 5.95
    yy = y + row * 1.15
    rect(s, x, yy, 5.6, 0.95, WHITE, line=LINE)
    rect(s, x, yy, 0.06, 0.95, BLUE, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.22, yy + 0.12, 2.0, 0.35, t, size=13, bold=True, color=NAVY)
    box(s, x + 0.22, yy + 0.5, 5.2, 0.35, d, size=10.5, color=MUTED)
box(s, M, y + 2.55, 11.6, 0.4, "每项 0–3 分，合计 15 分；水平分为 起步 / 初步 / 较好 / 成熟 四档。",
    size=13.5, color=INK)

# ---------------------------------------------------------------- 14 评分信度
s = new_slide(); y = header(s, "评分信度：把「AI 打分不可信」变成可测量的指标", "RELIABILITY", 14)
flow = [("第一次评分", "temperature = 0.0"), ("第二次评分", "temperature = 0.5"),
        ("逐维比较", "差值 0 / 1 / ≥2"), ("标记复核", "低一致性 → 人工复核")]
for i, (t, d) in enumerate(flow):
    x = M + i * 2.95
    rect(s, x, y + 0.2, 2.65, 1.25, WHITE, line=LINE)
    box(s, x + 0.16, y + 0.38, 2.3, 0.35, t, size=13, bold=True, color=NAVY)
    box(s, x + 0.16, y + 0.8, 2.3, 0.35, d, size=10.5, color=MUTED)
    if i < 3:
        box(s, x + 2.65, y + 0.7, 0.3, 0.3, "→", size=15, color=BLUE, align=PP_ALIGN.CENTER)
rect(s, M, y + 1.85, 11.6, 1.3, LIGHT)
box(s, M + 0.35, y + 2.05, 10.9, 0.9,
    "一致性「高 / 中 / 低」直接显示给学生与教师；低一致性不静默通过，而是提示复核。\n"
    "这是本项目在「评估」上的核心主张：可信度必须被报告，而不是被假设。",
    size=13.5, color=INK, spacing=1.4)

# ---------------------------------------------------------------- 15 用户反馈（虚拟）
s = new_slide(); y = header(s, "用户反馈（虚拟示例 · 待实测替换）", "FEEDBACK", 15)
quotes = [
    ("以前问 AI 它直接把答案给我；现在它反问我，我反而得自己想。", "学生 A"),
    ("核对史料那一步我才发现，AI 说的跟课本不一样，不能全信。", "学生 B"),
    ("自评时我觉得自己写得挺好，结果 AI 只给 2 分，差得有点多。", "学生 C"),
    ("这个能不能用到别的科目？", "学生 D"),
    ("以前最花时间是写人设提示词，现在点一下就能生成。", "教师 E"),
]
for i, (q, who) in enumerate(quotes):
    col = i % 2
    row = i // 2
    x = M + col * 5.95
    yy = y + row * 1.18
    w = 5.6 if i < 4 else 11.6
    rect(s, x, yy, w, 0.98, WHITE, line=LINE)
    box(s, x + 0.22, yy + 0.12, w - 1.5, 0.7, "「" + q + "」", size=11.5, color=INK, spacing=1.3)
    box(s, x + w - 1.3, yy + 0.32, 1.1, 0.3, who, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
box(s, M, y + 3.7, 11.6, 0.35, "⚠️ 以上为虚拟示例，仅用于演示结论框架；实际测试后将整条替换。",
    size=11.5, color=AMBER, italic=True)

# ---------------------------------------------------------------- 16 反馈分析框架
s = new_slide(); y = header(s, "反馈如何转成研究结论：一个四类编码框架", "ANALYSIS", 16)
codes = [("认知负荷", "「更累了，但我也更投入了」", "任务是否有挑战但可完成"),
         ("信任校准", "「原来 AI 也会说错」", "是否形成对 AI 的合理怀疑"),
         ("元认知", "「我以为我写得挺好」", "自评与 AI 评分的差距"),
         ("迁移意愿", "「能不能用到别的科目」", "能力是否被视为可迁移")]
rect(s, M, y, 11.6, 0.4, NAVY, shape=MSO_SHAPE.RECTANGLE)
for i, h in enumerate(["编码类别", "典型表述", "它想检验什么"]):
    box(s, M + 0.18 + [0, 2.6, 7.0][i], y + 0.1, [2.4, 4.2, 4.2][i], 0.3,
        h, size=11.5, bold=True, color=WHITE)
cy = y + 0.55
for k, (a, b, c) in enumerate(codes):
    if k % 2 == 0:
        rect(s, M, cy - 0.06, 11.6, 0.72, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.18, cy, 2.4, 0.6, a, size=12.5, bold=True, color=NAVY)
    box(s, M + 2.78, cy, 4.1, 0.6, b, size=11.5, color=MUTED)
    box(s, M + 7.18, cy, 4.1, 0.6, c, size=11.5, color=INK)
    cy += 0.78
box(s, M, cy + 0.15, 11.6, 0.45, "同一类编码在多次测试中重复出现，才上升为「研究结论」。",
    size=13.5, italic=True, color=BLUE)

# ---------------------------------------------------------------- 17 结论
s = new_slide(); y = header(s, "由反馈推出的四条待验证假设", "HYPOTHESES", 17)
hyps = [
    ("H1", "认知卸载被逆转", "从「要答案」转为「被逼着论证」，主观上更累、更投入。"),
    ("H2", "批判性 AI 素养可教", "核对史料能触发「AI 会错」的认知，并可能迁移到其他信源。"),
    ("H3", "元认知可校准", "自评与 AI 评分的差距，可作为元认知准确性的可测指标。"),
    ("H4", "教师门槛可降低", "一键生成限制词，显著降低自定义人物的备课成本。"),
]
cy = y + 0.15
for code, t, d in hyps:
    rect(s, M, cy, 11.6, 1.05, WHITE, line=LINE)
    rect(s, M, cy, 0.07, 1.05, TEAL, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.3, cy + 0.22, 0.7, 0.5, code, size=16, bold=True, color=TEAL)
    box(s, M + 1.1, cy + 0.18, 3.0, 0.4, t, size=14, bold=True, color=NAVY)
    box(s, M + 4.2, cy + 0.2, 7.1, 0.7, d, size=11.5, color=INK, spacing=1.3)
    cy += 1.22
box(s, M, cy + 0.05, 11.6, 0.4, "注意：这是假设，不是结论；需要实测数据检验。", size=11.5, color=AMBER, italic=True)

# ---------------------------------------------------------------- 18 局限
s = new_slide(); y = header(s, "局限", "LIMITATIONS", 18)
bullets(s, M, y, 11.4, [
    ("样本小、周期短。", "目前只有设计与小规模试用，任何「学习增益」都只是信号，不是证据。"),
    ("评分主观且不稳定。", "五维 rubric 只能近似历史思维；双评分能暴露不一致，但不能消除它。"),
    ("史实依赖史料库。", "史料库薄，练习就弱；冷门人物尤其明显。"),
    ("数据口径不一。", "不同调查对「使用 AI」的定义不同（16%–85%），引用时必须注明来源与样本。"),
], size=14.5, gap=0.15)

# ---------------------------------------------------------------- 19 未来工作
s = new_slide(); y = header(s, "未来工作", "FUTURE WORK", 19)
bullets(s, M, y, 11.4, [
    ("角色互换", "学生扮演历史人物，AI 扮演追问者——把「被质询」变成「被追问」。"),
    ("双人对质", "同场景引入两个立场相反的人物，让学生处理互相矛盾的叙述。"),
    ("更大规模的对照实验", "实验组 vs 常规 AI 答疑，测量历史思维与 AI 素养的变化。"),
    ("人物内容三语化", "让中文、繁体、英文课堂都能直接使用。"),
], size=14.5, gap=0.14)

# ---------------------------------------------------------------- 20 Q&A
s = new_slide()
rect(s, 0, 0, SW, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.2, SH, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 2.6, 11, 0.9, "Q & A", size=40, bold=True, color=WHITE)
box(s, M, 3.7, 11, 0.6, "谢谢 —— 欢迎提问与批评。", size=18, color=RGBColor(0xC7, 0xD3, 0xE4))
box(s, M, 5.4, 11, 0.4, "演示与代码：github.com/takakun09/ai-history-simulator", size=12,
    color=RGBColor(0x8F, 0x9F, 0xB5))
page_number(s, 20)

prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides._sldIdLst))
