# -*- coding: utf-8 -*-
"""Generate the course-project presentation (English)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

OUT = "/Users/rogan09/Documents/Codex/2026-10-09/new-chat/outputs/History_Simulator_Presentation.pptx"
AV = "/Users/rogan09/Documents/Codex/2026-10-09/new-chat/work/render"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x3A, 0x6E, 0xA5)
TEAL = RGBColor(0x2F, 0x6F, 0x5E)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x6B, 0x76, 0x86)
LINE = RGBColor(0xD8, 0xDE, 0xE7)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
M = 0.85  # margin


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def box(slide, x, y, w, h, text, size=17, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, italic=False, font=FONT, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        f = r.font
        f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sh.adjustments[0] = radius
        except Exception: pass
    return sh


def slide_number(slide, n):
    box(slide, SW - M - 1.0, SH - 0.62, 1.0, 0.3, str(n), size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def header(slide, title, kicker=None):
    rect(slide, 0, 0, 0.16, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
    y = 0.62
    if kicker:
        box(slide, M, y, 11, 0.3, kicker.upper(), size=11, bold=True, color=BLUE)
        y += 0.34
    box(slide, M, y, 11.6, 0.95, title, size=32, bold=True, color=NAVY, spacing=0.95)
    rect(slide, M, y + 1.02, 1.1, 0.045, BLUE, shape=MSO_SHAPE.RECTANGLE)
    return y + 1.35


def bullets(slide, x, y, w, items, size=17, gap=0.14, color=INK, bullet_color=BLUE):
    cy = y
    for it in items:
        if isinstance(it, tuple):
            head, sub = it
        else:
            head, sub = it, None
        rect(slide, x, cy + 0.085, 0.09, 0.09, bullet_color, shape=MSO_SHAPE.OVAL)
        tb = box(slide, x + 0.28, cy, w - 0.28, 0.4, head, size=size, color=color, spacing=1.05)
        n = 1
        if sub:
            box(slide, x + 0.28, cy + 0.33 * n + 0.02, w - 0.28, 0.6, sub, size=size - 3.5,
                color=MUTED, spacing=1.15)
            n += 1.6
        cy += 0.33 * n + gap + 0.22
    return cy


# ---------------------------------------------------------------- 1 title
s = blank()
rect(s, 0, 0, SW, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.22, SH, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 2.05, 10.6, 0.4, "AI IN EDUCATION  ·  COURSE PROJECT", size=13, bold=True, color=RGBColor(0x8F,0xB6,0xE0))
box(s, M, 2.6, 11.2, 1.9, "An AI That Might Be Wrong", size=54, bold=True, color=WHITE, spacing=0.95)
box(s, M, 4.35, 10.4, 1.0,
    "Teaching historical thinking by making students cross-check a biased historical narrator",
    size=20, color=RGBColor(0xC7,0xD3,0xE4), spacing=1.15)
rect(s, M, 5.6, 1.4, 0.04, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 5.85, 11, 0.4, "[Team members]   ·   October 2026", size=14, color=RGBColor(0x9F,0xAE,0xC4))
slide_number(s, 1)

# ---------------------------------------------------------------- 2 problem
s = blank(); y = header(s, "Students ask AI for answers, not judgment", "The problem")
bullets(s, M, y, 11.3, [
    "Generative AI makes a confident answer available in seconds — and makes it easy to skip thinking.",
    "Large language models hallucinate and flatter the user. Students rarely check.",
    "The abilities history actually requires — sourcing, contextualizing, corroborating — are not taught by telling.",
], size=20)
rect(s, M, 5.5, 11.6, 1.25, LIGHT)
box(s, M + 0.45, 5.75, 10.7, 0.8,
    "The more useful the tool feels, the less the student has to think. That is the design problem we started from.",
    size=19, italic=True, color=NAVY, spacing=1.15)
slide_number(s, 2)

# ---------------------------------------------------------------- 3 gap
s = blank(); y = header(s, "What is missing from the research", "Research gap")
bullets(s, M, y, 11.3, [
    ("AI role-play is well studied for language learning — far less for disciplinary reasoning.",
     "Most systems are built to converse fluently, not to train the evaluation of evidence."),
    ("Historical thinking is usually assessed by essay, rarely by an interactive system.",
     "There is no widely used instrument for measuring it inside a dialogue."),
    ("Almost no work treats the model's unreliability as designed teaching material.",
     "The bias is treated as a defect to remove, rather than a source of cognitive conflict."),
], size=19, gap=0.1)
slide_number(s, 3)

# ---------------------------------------------------------------- 4 idea
s = blank(); y = header(s, "Make the AI a witness, not an encyclopedia", "The idea")
box(s, M, y, 6.5, 0.4, "The model plays a historical actor with a stake and a point of view.", size=19, color=INK)
bullets(s, M, y + 0.6, 6.6, [
    "It is allowed to be partial, self-serving, and sometimes simply wrong.",
    "The student interrogates it, then checks its claims against primary sources.",
    "The student decides, and has to defend the decision with evidence.",
], size=17)
rect(s, 7.5, 1.98, 5.4, 3.55, WHITE, line=LINE)
for i, (fn, label) in enumerate([("zhuge-liang.png","Zhuge Liang"), ("socrates.png","Socrates"), ("lincoln.png","Abraham Lincoln")]):
    x = 7.72 + i * 1.68
    p = os.path.join(AV, fn)
    if os.path.exists(p):
        s.shapes.add_picture(p, Inches(x), Inches(2.28), height=Inches(1.7))
    box(s, x - 0.08, 4.08, 1.86, 0.35, label, size=12, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 7.85, 4.65, 4.7, 0.9, "Three figures, each with a documented position and a set of primary sources.",
    size=13.5, color=MUTED, spacing=1.15)
slide_number(s, 4)

# ---------------------------------------------------------------- 5 objectives
s = blank(); y = header(s, "What students should leave with", "Objectives")
bullets(s, M, y, 11.3, [
    ("Understand a historical decision from more than one vantage point.",
     "Not a single correct narrative, but the constraint each actor was under."),
    ("Source, contextualize, and corroborate.",
     "The three operations that define historical thinking in the literature."),
    ("Explain why people acted as they did — without excusing them.",
     "Historical empathy as perspective recognition."),
    ("Defend a decision with evidence.",
     "The output is a written argument, not a correct answer."),
    ("Treat a confident model as a source to be checked.",
     "AI literacy that transfers beyond the history classroom."),
], size=15.5, gap=0.0)
slide_number(s, 5)

# ---------------------------------------------------------------- 6 theory
s = blank(); y = header(s, "Theoretical grounding", "Foundations")
rows = [
    ("Historical thinking", "Wineburg", "Sourcing, contextualization, corroboration", "Rubric dimensions 1–3"),
    ("Historical reasoning", "van Drie & van Boxtel", "Asking questions, using sources, arguing", "Decision task design"),
    ("Situated learning", "Lave & Wenger", "Learning inside an authentic activity", "Role-play format"),
    ("Productive failure", "Kapur", "Struggle before instruction", "AI withholds the answer"),
]
rect(s, M, y + 0.05, 11.6, 0.42, NAVY, shape=MSO_SHAPE.RECTANGLE)
for i, h in enumerate(["Framework", "Source", "What it says", "Where it lands in the build"]):
    box(s, M + 0.18 + [0, 2.5, 5.0, 8.6][i], y + 0.13, [2.4, 2.5, 3.6, 3.0][i], 0.3, h, size=12, bold=True, color=WHITE)
cy = y + 0.62
for k, (a, b, c, d) in enumerate(rows):
    if k % 2 == 0:
        rect(s, M, cy - 0.06, 11.6, 0.72, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.18, cy, 2.4, 0.6, a, size=13, bold=True, color=NAVY)
    box(s, M + 2.68, cy, 2.4, 0.6, b, size=12.5, color=MUTED)
    box(s, M + 5.18, cy, 3.4, 0.6, c, size=12.5, color=INK)
    box(s, M + 8.78, cy, 2.9, 0.6, d, size=12.5, color=BLUE)
    cy += 0.78
box(s, M, cy + 0.05, 11.6, 0.5, "Every framework maps to a concrete feature or rubric row — none of them is decoration.",
    size=15, italic=True, color=NAVY)
slide_number(s, 6)

# ---------------------------------------------------------------- 7 architecture
s = blank(); y = header(s, "System design", "Architecture")
mods = [("Scenario engine", "Sets the moment and the decision points"), ("Source library", "Primary and secondary sources, with provenance"),
        ("Persona agent", "Plays the historical actor, grounded in sources"), ("Diagnosis & teacher console", "Scores arguments, aggregates the class profile")]
cy = y + 0.15
for i, (t, d) in enumerate(mods):
    x = M + i * 2.98
    rect(s, x, cy, 2.72, 1.5, WHITE, line=LINE)
    rect(s, x, cy, 2.72, 0.075, BLUE, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.22, cy + 0.28, 2.3, 0.4, t, size=15, bold=True, color=NAVY, spacing=1.05)
    box(s, x + 0.22, cy + 0.78, 2.3, 0.6, d, size=11.5, color=MUTED, spacing=1.12)
box(s, M, cy + 1.85, 11.6, 0.4, "Student", size=13, bold=True, color=BLUE)
rect(s, M, cy + 2.2, 11.6, 0.05, LINE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, cy + 2.45, 11.6, 1.3,
    "Built entirely on a pretrained model accessed through an API, plus prompt engineering and a small source\n"
    "library. No model was trained or fine-tuned — a deliberate scope decision, not a limitation we stumbled into.",
    size=15, color=INK, spacing=1.25)
slide_number(s, 7)

# ---------------------------------------------------------------- 8 loop
s = blank(); y = header(s, "The learning loop", "How a session runs")
steps = [("Interrogate", "Ask the historical actor"), ("Check sources", "Judge the claim against evidence"),
         ("Decide", "Choose at a decision point"), ("Justify", "Write the argument"), ("Feedback", "Scores, agreement, profile")]
bw, gap = 2.12, 0.24
for i, (t, d) in enumerate(steps):
    x = M + i * (bw + gap)
    rect(s, x, y + 0.35, bw, 1.45, WHITE, line=LINE)
    rect(s, x, y + 0.35, bw, 0.075, TEAL, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.16, y + 0.55, bw - 0.32, 0.35, t, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    box(s, x + 0.16, y + 0.98, bw - 0.32, 0.7, d, size=11, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.1)
    if i < len(steps) - 1:
        box(s, x + bw, y + 0.95, gap, 0.3, "→", size=16, color=BLUE, align=PP_ALIGN.CENTER)
box(s, M, y + 2.15, 11.6, 1.5,
    "The loop is the pedagogical claim. Students are never handed the answer: the model is a source to be\n"
    "cross-examined, and the gradeable artefact is their reasoning, not their recall.",
    size=17, color=INK, spacing=1.3)
slide_number(s, 8)

# ---------------------------------------------------------------- 9 build
s = blank(); y = header(s, "What we built", "Working prototype")
cards = [("Student console", ["Pick a figure, then chat", "One-click source check", "Decision + written argument", "Personal profile"]),
         ("Teacher console", ["Author figures and prompts", "Edit the source list", "Design decision points", "Publish to students"]),
         ("Class analytics", ["Submissions and checks", "Five-dimension averages", "Weakest dimension", "Recent records"])]
for i, (t, items) in enumerate(cards):
    x = M + i * 4.02
    rect(s, x, y, 3.72, 2.9, WHITE, line=LINE)
    rect(s, x, y, 3.72, 0.075, [BLUE, TEAL, NAVY][i], shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.28, y + 0.32, 3.2, 0.4, t, size=17, bold=True, color=NAVY)
    cy2 = y + 0.95
    for it in items:
        box(s, x + 0.28, cy2, 3.2, 0.35, "·  " + it, size=12.5, color=INK)
        cy2 += 0.44
box(s, M, y + 3.25, 11.6, 0.6,
    "Screenshots to insert: student chat · source check · decision scoring · profile · teacher editor · analytics.",
    size=13, italic=True, color=MUTED)
slide_number(s, 9)

# ---------------------------------------------------------------- 10 assessment
s = blank(); y = header(s, "Assessing historical thinking", "Assessment design")
dims = [("Sourcing", "Notes the origin and position of a claim"), ("Contextualization", "Places a decision in its period"),
        ("Corroboration", "Compares more than one source"), ("Argumentation", "Claim — evidence — reasoning"),
        ("Historical empathy", "Explains motive under constraint")]
box(s, M, y, 7.4, 0.35, "Five dimensions, each scored 0–3", size=16, bold=True, color=NAVY)
cy = y + 0.55
for i, (a, b) in enumerate(dims):
    if i % 2 == 0:
        rect(s, M, cy - 0.05, 7.4, 0.62, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.16, cy + 0.06, 2.2, 0.4, a, size=13.5, bold=True, color=NAVY)
    box(s, M + 2.45, cy + 0.06, 4.8, 0.5, b, size=13, color=MUTED)
    cy += 0.68
rect(s, 8.75, y, 3.7, 3.9, WHITE, line=LINE)
box(s, 9.05, y + 0.3, 3.1, 0.4, "Reliability", size=16, bold=True, color=BLUE)
box(s, 9.05, y + 0.85, 3.1, 2.7,
    "Scoring runs twice, at two temperatures.\n\nThe two passes are averaged, and each dimension is tagged high / medium / low agreement.\n\nLow agreement flags the case for human review instead of silently trusting the score.",
    size=12.5, color=INK, spacing=1.25)
slide_number(s, 10)

# ---------------------------------------------------------------- 11 value
s = blank(); y = header(s, "Educational value", "Why it matters")
bullets(s, M, y, 11.3, [
    ("It turns the most convenient feature of AI into the object of study.",
     "Instead of blocking the model, students learn to interrogate it."),
    ("It makes an invisible disciplinary skill visible and practiceable.",
     "Sourcing and corroboration become actions, not essay conventions."),
    ("It gives teachers a class-level view of reasoning, not just scores.",
     "Which dimension is weak for this cohort, and which figure provoked it."),
    ("It transfers outside history.",
     "Treating a fluent, confident source with suspicion is a general competence."),
], size=18, gap=0.08)
slide_number(s, 11)

# ---------------------------------------------------------------- 12 feasibility
s = blank(); y = header(s, "Feasibility", "What we can already claim")
bullets(s, M, y, 6.6, [
    "A working prototype exists — student, teacher, and analytics consoles.",
    "No model training; a pretrained model behind an API plus prompt engineering.",
    "Runs on a laptop with the Python standard library; no dependencies to install.",
    "An offline mode lets the demonstration run without any API key.",
], size=16, gap=0.1)
rect(s, 7.75, y, 4.7, 3.4, LIGHT)
box(s, 8.05, y + 0.3, 4.1, 0.4, "What is still missing", size=15, bold=True, color=NAVY)
box(s, 8.05, y + 0.85, 4.1, 2.3,
    "Evidence of learning.\n\nWe have a design and a working artefact. We do not yet have pre/post data — a short qualitative pilot is the next step.",
    size=13, color=INK, spacing=1.3)
slide_number(s, 12)

# ---------------------------------------------------------------- 13 limitations
s = blank(); y = header(s, "Limitations", "Honest constraints")
bullets(s, M, y, 11.3, [
    "The construct is hard to measure: five rubric dimensions can only approximate historical thinking.",
    "Automated scoring of open-ended argument is unstable; agreement is reported rather than hidden.",
    "The pilot is small and short. Any learning gain we report is a signal, not proof.",
    "Historical accuracy depends on the source library; a thin library produces a weak exercise.",
    "Sensitive topics need curation — the same realism that makes the design work can also misfire.",
], size=17, gap=0.08)
slide_number(s, 13)

# ---------------------------------------------------------------- 14 ethics
s = blank(); y = header(s, "Ethical and social implications", "Responsibility")
bullets(s, M, y, 11.3, [
    ("Designed bias must be labelled as design.",
     "Students must not conclude that the past is simply unknowable."),
    ("Student dialogue is sensitive data.",
     "It is stored locally; any classroom deployment needs consent and a retention policy."),
    ("Cost and language shape who can use it.",
     "Hosted APIs cost money and carry a Western bias; a local model path matters for equity."),
    ("The teacher is augmented, not replaced.",
     "The system surfaces reasoning; it does not make the judgement."),
], size=17, gap=0.06)
slide_number(s, 14)

# ---------------------------------------------------------------- 15 decisions
s = blank(); y = header(s, "Critical design decisions", "Trade-offs we chose")
rows = [
    ("Let the AI be biased", "Historical thinking can only be practised against a partial source", "Risk of reinforcing error → labelled, curated scenarios"),
    ("Ground every reply in sources", "Keeps the bias designed and gradeable rather than random", "Slower to author; a thin library weakens the task"),
    ("Score twice and report agreement", "Makes scoring reliability visible", "Roughly doubles scoring cost → opt-in"),
    ("No fine-tuning", "Two-week scope; prompts plus retrieval are enough", "Ceiling on stylistic control; acceptable here"),
]
cy = y + 0.05
for i, (a, b, c) in enumerate(rows):
    if i % 2 == 0:
        rect(s, M, cy - 0.06, 11.6, 0.95, LIGHT, shape=MSO_SHAPE.RECTANGLE)
    box(s, M + 0.18, cy + 0.06, 2.9, 0.6, a, size=14, bold=True, color=NAVY)
    box(s, M + 3.25, cy + 0.06, 3.9, 0.8, b, size=12.5, color=INK)
    box(s, M + 7.35, cy + 0.06, 4.1, 0.8, c, size=12.5, color=MUTED)
    cy += 1.02
box(s, M, cy + 0.05, 11.6, 0.4, "left: decision      middle: why      right: the cost we accepted",
    size=11.5, italic=True, color=MUTED)
slide_number(s, 15)

# ---------------------------------------------------------------- 16 future
s = blank(); y = header(s, "Research gaps and future work", "Where this goes next")
bullets(s, M, y, 11.3, [
    ("A validated instrument for critical AI literacy.",
     "Our rubric is a first draft; it needs reliability and validity evidence across cohorts."),
    ("Longer-term transfer.",
     "Does source-checking inside the simulation change how students read sources outside it?"),
    ("Multimodal scenarios.",
     "Maps, images, and objects as additional conflicting evidence."),
    ("Localisation.",
     "Non-Western figures and non-English sources, using locally hosted models."),
], size=17, gap=0.04)
slide_number(s, 16)

# ---------------------------------------------------------------- 17 end
s = blank()
rect(s, 0, 0, SW, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.22, SH, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, M, 2.5, 11, 1.2, "Questions", size=44, bold=True, color=WHITE)
box(s, M, 3.85, 11, 0.6, "The prototype is running on the laptop — we can walk through it live.", size=18, color=RGBColor(0xC7,0xD3,0xE4))
box(s, M, 5.6, 11, 0.4, "[Team members]   ·   AI in Education   ·   October 2026", size=13, color=RGBColor(0x8F,0x9F,0xB5))
slide_number(s, 17)

prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
